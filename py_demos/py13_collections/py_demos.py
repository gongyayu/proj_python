import streamlit as st
from py_mod import line_print, st_print, st_code, outlines

### StartofFunc###


def py_polinomials():
    total = 0
    num = 8
    for x in range(num):
        for y in range(x, num):
            total += x*y
            st_code(f"x: {x}, y: {y}, total: {total}")
            if total > 30:
                st_code(f"inner total: {total}")
                break
    st_code(f"final total: {total}")
### EndofCodeSection###


def py_translator():
    import asyncio
    from googletrans import Translator

    async def main():
        translator = Translator()
        txt = st.text_input("Enter English to be translated into Chinese: ")
        # txt = "Our team is developing a useful app for people worldwide"
        if txt:
            output = await translator.translate(txt, dest="zh-cn")
            st_print(f"{txt} -> {output.text}")
    asyncio.run(main())
### EndofCodeSection###


def pymsword():
    from docx import Document
    doc = Document('msword.docx')
    for p in doc.paragraphs:
        print(p.text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                print(cell.text, end="\t")
            print()
### EndofCodeSection###


def py_create_triangle():
    r = 10
    for x in range(r):
        line_print(f"' '*{(r-x-1)}+'*'*{(2*x+1)}")
    st_print(line_print(f''))
### EndofCodeSection###


def py_screenshot():
    import pyautogui
    screenshot = pyautogui.screenshot()
    screenshot.save('screenshot.png')
    st_print("Screenshot saved as 'screenshot.png'")
### EndofCodeSection###


def pyturtle():
    import turtle

    # Set up the turtle
    t = turtle.Turtle()
    t.speed(1)

    # Draw a square
    for _ in range(4):
        t.forward(100)
        t.right(90)

    # Hide the turtle and display the window
    t.hideturtle()
    turtle.done()
### EndofCodeSection###
